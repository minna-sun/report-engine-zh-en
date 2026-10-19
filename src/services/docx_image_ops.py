"""图片插入 Mixin（v5.0 · 由 docx_engine.py 拆分而来）。

职责：``DocxEngine`` 的**图片插入**能力——在锚点段落后（或锚点所在表格单元格内）
新建段落并插入图片，锁定纵横比并居中。

以 Mixin 形式被 ``DocxEngine`` 继承。

依赖：``services.docx_helpers``（底层纯函数）、``python-docx``。
"""
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm

from common.logger import warning

from services.docx_helpers import (
    _is_in_table,
    _new_paragraph,
    _paragraph_from_element,
)
from services.ooxml_service import _q


class ImageOpsMixin:
    """图片插入 Mixin。"""

    def insert_image(self, doc: Document, anchor: object,
                     image_data: bytes, width_cm: float) -> bool:
        """锚点段落后新建段落并插入图片，锁定纵横比、居中。

        若锚点段落在表格单元格内，则在该单元格内追加段落插入图片；
        否则在锚点段落之后、正文中插入新段落。
        """
        try:
            anchor_para = anchor
            warning(f"[docx_engine] insert_image anchor text: "
                    f"{(anchor.text or '').strip()!r}, is_in_table: {_is_in_table(anchor)}")
            if _is_in_table(anchor_para):
                # ── 锚点在表格内：找到对应的 Cell 对象，在其内部追加段落 ──
                p = anchor_para._p
                tc = p.getparent()
                while tc is not None and tc.tag != _q("tc"):
                    tc = tc.getparent()
                cell_obj = None
                if tc is not None:
                    for tbl in doc.tables:
                        for row in tbl.rows:
                            for cell in row.cells:
                                if cell._tc is tc:
                                    cell_obj = cell
                                    break
                            if cell_obj is not None:
                                break
                        if cell_obj is not None:
                            break
                if cell_obj is not None:
                    warning(f"[docx_engine] insert_image：锚点在表格内，"
                            f"图片追加至单元格（width={width_cm}cm）")
                    new_para = cell_obj.add_paragraph()
                    new_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = new_para.add_run()
                    run.add_picture(BytesIO(image_data), width=Cm(width_cm))
                    # ── 行高改为 atLeast，让行随图片高度自动增长 ──
                    try:
                        _tr = tc.getparent() if tc is not None else None
                        if _tr is not None and _tr.tag == _q("tr"):
                            _trPr = _tr.find(_q("trPr"))
                            if _trPr is None:
                                _trPr = OxmlElement("w:trPr")
                                _tr.insert(0, _trPr)
                            _trH = _trPr.find(_q("trHeight"))
                            if _trH is None:
                                _trH = OxmlElement("w:trHeight")
                                _trH.set(_q("val"), "0")
                                _trPr.append(_trH)
                            _trH.set(_q("hRule"), "atLeast")
                    except Exception as _e:
                        warning(f"[docx_engine] insert_image 行高调整失败（不影响插图）：{_e}")
                    return True
                # 找不到 Cell 对象，回退到正文插入路径
                warning("[docx_engine] insert_image：锚点在表格内但未找到 Cell 对象，"
                        "回退正文段落插入")
            else:
                warning(f"[docx_engine] insert_image：锚点在正文，"
                        f"插入到锚点后新段落（width={width_cm}cm）")
            # ── 锚点在正文（或回退路径）：锚点后插入新段落 ──
            new_p = _new_paragraph(doc)
            anchor_para._p.addnext(new_p)
            para = _paragraph_from_element(new_p, doc)
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run()
            run.add_picture(BytesIO(image_data), width=Cm(width_cm))
            return True
        except Exception as e:
            warning(f"[docx_engine] insert_image 失败：{e}")
            return False
