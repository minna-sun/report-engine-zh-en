"""导航 / 定位 Mixin（v5.0 · 由 docx_engine.py 拆分而来）。

职责：``DocxEngine`` 的**文档定位**能力——章节查找、锚点查找、占位符查找、
表格定位、章节范围计算、分页符检测，以及章节定位缓存的失效控制。

以 Mixin 形式被 ``DocxEngine`` 继承；内部通过 ``self`` 调用，
依赖宿主类提供 ``self._section_cache``（由 ``DocxEngine.__init__`` 初始化）。

依赖：``services.docx_helpers``（底层纯函数）、``services.ooxml_service``。
"""
import re as _re
from typing import Optional

from docx import Document

from common.logger import warning

from services.docx_helpers import (
    _element_after,
    _has_page_break_before,
    _is_chapter_title_para,
    _is_in_table,
    _iter_all_paragraphs,
    _normalize_colon,
    _para_is_bold,
)


class NavigationMixin:
    """文档导航 / 定位 Mixin。"""

    def invalidate_section_cache(self, doc: Optional[Document] = None) -> None:
        """清除章节定位缓存。

        Args:
            doc: 仅清除该文档的缓存项；传 ``None`` 则清空全部缓存。
        """
        if doc is None:
            self._section_cache.clear()
        else:
            doc_id = id(doc)
            keys = [k for k in self._section_cache if k[0] == doc_id]
            for k in keys:
                del self._section_cache[k]

    def find_section(self, doc: Document, section_name: str,
                      require_bold: bool = True) -> Optional[object]:
        """匹配文本含 section_name 的段落；默认要求加粗且忽略表格内段落。

        结果会被缓存（以 doc 对象 id + 参数为 key），文档结构改变后需调用
        ``invalidate_section_cache`` 使缓存失效。
        """
        cache_key = (id(doc), section_name, require_bold)
        if cache_key in self._section_cache:
            cached = self._section_cache[cache_key]
            warning(f"[docx_engine] find_section：缓存命中 {section_name!r} → "
                    f"{'找到' if cached is not None else '未找到'}")
            return cached

        result = None
        for para in doc.paragraphs:
            if _is_in_table(para):
                continue
            if require_bold and not _para_is_bold(para):
                continue
            _needle = _re.sub(r"\s+", "", section_name)
            if _needle in _re.sub(r"\s+", "", para.text or ""):
                result = para
                break

        if result is not None:
            warning(f"[docx_engine] find_section：找到 {section_name!r} "
                    f"→ {(result.text or '').strip()!r}")
        else:
            warning(f"[docx_engine] find_section：未找到 {section_name!r}")

        self._section_cache[cache_key] = result
        return result

    def find_paragraph(self, doc: Document, section: object,
                       anchor_text: str) -> Optional[object]:
        """在 section 段落后、下一个加粗标题前，查找含 anchor_text 的段落（子串匹配）。

        注意：python-docx 每次 ``doc.paragraphs`` 都返回新的 Paragraph 包装对象，
        故以底层 ``._p`` 元素身份比较，而非包装对象身份。
        """
        section_name = (section.text or "").strip() if section is not None else "<none>"

        started = False
        for para in doc.paragraphs:
            if para._p is section._p:
                started = True
                continue
            if started:
                if _para_is_bold(para) and para.text.strip():
                    break
                _para_text = para.text or ""
                if _normalize_colon(anchor_text) in _normalize_colon(_para_text):
                    warning(f"[docx_engine] find_paragraph：命中（正文段落）"
                            f" section={section_name!r} anchor={anchor_text!r}"
                            f" text={_para_text[:40]!r}")
                    return para

        # 回退 1：搜索表格单元格内段落。
        # 真实模板（如 Sample Description 的 Material Number / Supplier 等字段）
        # 常把「标签|值」放在表格单元格里，而 python-docx 的 doc.paragraphs
        # 只含正文顶层段落、不含表格内段落，故必须回退搜索，否则字段永远填不上。
        if section is not None and started:
            # 章节边界：section 之后第一个加粗正文标题
            next_title = None
            for para in doc.paragraphs:
                if para._p is section._p:
                    continue
                if (_element_after(section._p, para._p)
                        and _para_is_bold(para) and para.text.strip()):
                    next_title = para
                    break
            # ── 诊断：打印章节内所有候选段落（仅 sample_title 相关章节）─────────
            if "样品" in section_name or "sample" in section_name.lower():
                warning(f"[find_paragraph] 诊断 section={section_name!r} "
                        f"next_title={((next_title.text or '') if next_title else None)!r}")
                for _dp in _iter_all_paragraphs(doc):
                    if not _element_after(section._p, _dp._p):
                        continue
                    if next_title is not None and _dp._p is next_title._p:
                        break
                    _dt = (_dp.text or "").strip()
                    if _dt:
                        warning(f"[find_paragraph] 候选段落: {_dt[:60]!r}")
            # ─────────────────────────────────────────────────────────────────
            for para in _iter_all_paragraphs(doc):
                if not _element_after(section._p, para._p):
                    continue
                if next_title is not None and para._p is next_title._p:
                    break
                if _normalize_colon(anchor_text) in _normalize_colon(para.text or ""):
                    warning(f"[docx_engine] find_paragraph：命中（表格单元格）"
                            f" section={section_name!r} anchor={anchor_text!r}"
                            f" text={((para.text or '')[:20])!r}")
                    return para

        # 回退 2：全文（兼容无后续标题 / section 未命中情况）
        if not started:
            for para in doc.paragraphs:
                if _normalize_colon(anchor_text) in _normalize_colon(para.text or ""):
                    warning(f"[docx_engine] find_paragraph：命中（全文兜底，section 未匹配）"
                            f" section={section_name!r} anchor={anchor_text!r}"
                            f" text={((para.text or '')[:20])!r}")
                    return para

        warning(f"[docx_engine] find_paragraph：未找到"
                f" section={section_name!r} anchor={anchor_text!r}")
        return None

    def find_placeholder(self, doc: Document, section: object,
                         placeholder: str) -> Optional[object]:
        """在 section 章节范围内精确匹配占位符段落。

        ``placeholder`` 应包含花括号（如 ``"{front}"``）。匹配条件：段落文本
        ``strip()`` 后与 ``placeholder`` 完全相等。先搜正文段落，再搜表格单元格。
        不做全文回退——找不到直接返回 ``None``，调用方可记录警告。
        """
        if section is None:
            warning(f"[docx_engine] find_placeholder：section 为 None，"
                    f"无法定位 {placeholder!r}")
            return None

        section_name = (section.text or "").strip()
        target = (placeholder or "").strip()

        # ── 正文段落（章节标题到下一个章节标题之间）────────────────────────
        started = False
        for para in doc.paragraphs:
            if para._p is section._p:
                started = True
                continue
            if started:
                if _para_is_bold(para) and para.text.strip():
                    break   # 到达下一个章节标题
                if (para.text or "").strip() == target:
                    warning(f"[docx_engine] find_placeholder：命中（正文段落）"
                            f" section={section_name!r} placeholder={placeholder!r}")
                    return para

        if not started:
            warning(f"[docx_engine] find_placeholder：章节 {section_name!r} 未找到，"
                    f"无法定位占位符 {placeholder!r}")
            return None

        # ── 表格单元格（同章节范围内）────────────────────────────────────────
        next_title = None
        for para in doc.paragraphs:
            if para._p is section._p:
                continue
            if (_element_after(section._p, para._p)
                    and _para_is_bold(para) and para.text.strip()):
                next_title = para
                break
        for para in _iter_all_paragraphs(doc):
            if not _element_after(section._p, para._p):
                continue
            if next_title is not None and para._p is next_title._p:
                break
            if (para.text or "").strip() == target:
                warning(f"[docx_engine] find_placeholder：命中（表格单元格）"
                        f" section={section_name!r} placeholder={placeholder!r}")
                return para

        warning(f"[docx_engine] find_placeholder：未找到"
                f" section={section_name!r} placeholder={placeholder!r}")
        return None

    def find_table(self, doc: Document, section: object,
                   anchor_text: str) -> Optional[object]:
        """查找 anchor_text 段落后第一个 Table 对象。"""
        anchor = self.find_paragraph(doc, section, anchor_text)
        if anchor is None:
            return None
        for tbl in doc.tables:
            if _element_after(anchor._p, tbl._tbl):
                return tbl
        return doc.tables[0] if doc.tables else None

    def list_sections(self, doc: Document):
        """返回文档中所有测试章节标题段落列表 ``[(文本, 段落句柄), ...]``（按文档顺序）。

        判定：加粗 + 非表格内 + 文本非空 + 不在黑名单（见 ``_BLACKLIST_TITLES``）。
        """
        out = []
        try:
            for para in doc.paragraphs:
                if _is_chapter_title_para(para):
                    out.append((para.text or "", para))
        except Exception as e:
            warning(f"[docx_engine] list_sections 失败：{e}")
        return out

    def find_paragraph_by_text(self, doc: Document, text: str,
                               exact: bool = False) -> Optional[object]:
        """全局搜索含 text（或文本等于 text）的段落，返回首个匹配句柄。

        搜索顺序：
          1. 普通段落（doc.paragraphs）——优先
          2. 所有表格单元格内段落——兜底（用于占位符在表格中的场景）
        """
        def _matches(t: str) -> bool:
            return (t.strip() == text.strip()) if exact else (text in t)

        try:
            # 1. 普通段落
            for para in doc.paragraphs:
                if _matches(para.text or ""):
                    return para
            # 2. 表格单元格内段落
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for para in cell.paragraphs:
                            if _matches(para.text or ""):
                                return para
        except Exception as e:
            warning(f"[docx_engine] find_paragraph_by_text 失败：{e}")
        return None

    def find_section_range(self, doc: Document, section_name: str):
        """返回章节的起止段落。边界判定：
          1. 下一个加粗非空段落（测试章节标题）
          2. 或下一个黑名单标题（如 "The End"、"声明" 等固定页）
        两者取先到者作为结束边界。
        """
        from services.docx_helpers import _is_blacklist_title
        start = self.find_section(doc, section_name)
        if start is None:
            return None, None
        started = False
        end = None
        for para in doc.paragraphs:
            if para._p is start._p:
                started = True
                continue
            if started:
                if _is_in_table(para):
                    continue
                # 边界1：加粗非空段落（普通章节标题）
                if _para_is_bold(para) and (para.text or "").strip():
                    end = para
                    break
                # 边界2：黑名单标题（The End / 声明 等固定页）
                if _is_blacklist_title(para):
                    end = para
                    break
        return start, end

    def has_page_break_before(self, para: object) -> bool:
        """判断段落前是否已存在分页符。"""
        return _has_page_break_before(para)

    def _find_anchor_in_section(self, doc: Document, section_name: str,
                                anchor_text: str) -> Optional[object]:
        """在 section_name 章节范围内搜索含 anchor_text 的段落。

        找到返回该段落；章节不存在或章节内找不到时返回 None 并记录警告。
        不执行任何全文回退搜索。
        """
        if not section_name:
            warning(f"[docx_engine] _find_anchor_in_section：未传入章节名")
            return None
        # 章节结构在 clone_section / delete_section 后可能已变化，清除缓存确保重新定位
        self.invalidate_section_cache(doc)
        sec = self.find_section(doc, section_name)
        if sec is None:
            warning(f"[docx_engine] _find_anchor_in_section：未找到章节 {section_name!r}")
            return None
        para = self.find_paragraph(doc, sec, anchor_text)
        if para is None:
            warning(f"[docx_engine] _find_anchor_in_section：章节 {section_name!r} "
                    f"内未找到锚点 {anchor_text!r}")
            return None
        return para

    @staticmethod
    def _first_table_after(doc: Document, anchor_para) -> Optional[object]:
        """返回 anchor_para 段落后第一个 Table 对象（anchor 为 None 时返回首个表格）。"""
        if anchor_para is None:
            return doc.tables[0] if doc.tables else None
        for tbl in doc.tables:
            if _element_after(anchor_para._p, tbl._tbl):
                return tbl
        return None
